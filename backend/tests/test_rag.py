import pytest

from tools import rag
from tools.rag import chunk_cache, paths, pipeline
from tools.rag.chunking import ChunkingMetrics, PreparedChunk
from tools.rag.models import BuildOptions, RagRequest
from tools.settings import RagSettings


def test_rag_searches_knowledge_memory_and_skills(rag_env):
    rag_env.write("knowledge", "api.md", "OpenAI compatible chat completions notes")
    rag_env.write("memory", "MEMORY.md", "User prefers compact agent prompts")
    rag_env.write("skill", "backend", "Use LangGraph assistant_step tool_call loop")

    results = rag.search("LangGraph compact prompts", rag_env.settings(max_results=5))

    joined = "\n".join(results)
    assert "sourceType: skill" in joined
    assert "sourceType: memory" in joined


def test_rag_auto_context_uses_same_search(rag_env):
    rag_env.write("memory", "MEMORY.md", "approvalRequired means preview only")

    context = rag.auto_context("approvalRequired", rag_env.settings(mode="on", max_results=1))

    assert context is not None
    assert context.startswith("autoRagResult:")
    assert "preview only" in context


def test_rag_disabled_blocks_model_search(tmp_path):
    settings = RagSettings(mode="off", memory_root=str(tmp_path))

    with pytest.raises(ValueError, match="disabled"):
        rag.search("anything", settings)


def test_rag_builds_a_persistent_hybrid_index(rag_env):
    rag_env.write("knowledge", "api.md", "vector search stores chunks")
    settings = rag_env.settings()

    status = rag.index_status(settings)
    results = rag.search("vector chunks", settings)

    assert status["index"] == "hybrid"
    assert status["retrievers"] == ["bm25", "vector"]
    assert status["embedding"] == "tfidf-char-ngram"
    assert status["chunk_count"] == 1
    assert (rag_env.root / "index" / "index.pkl").is_file()
    assert "sourceType: knowledge" in "\n".join(results)


def test_rag_replaces_incompatible_index_without_unpickling_it(rag_env):
    rag_env.write("knowledge", "note.md", "新向量索引")
    index_path = rag_env.root / "index" / "index.pkl"
    index_path.parent.mkdir()
    index_path.write_bytes(b"incompatible dense index that must not be loaded")

    status = rag.index_status(rag_env.settings())

    assert status["embedding"] == "tfidf-char-ngram"
    assert index_path.read_bytes().startswith(rag.INDEX_MAGIC)


def test_rag_llm_chunk_cache_only_reprocesses_changed_documents(rag_env, monkeypatch):
    first = rag_env.write("knowledge", "first.md", "first version")
    second = rag_env.write("knowledge", "second.md", "unchanged")
    calls = []

    def fake_split(text, **kwargs):
        calls.append(text)
        return [PreparedChunk(text, "planned")], ChunkingMetrics(llm_calls=1)

    monkeypatch.setattr(pipeline, "split_document", fake_split)
    settings = rag_env.settings()
    options = BuildOptions(split_mode="llm", chunk_model="chunk-model")

    initial = rag.index_status(settings, options)
    first.write_text("first version updated", encoding="utf-8")
    updated = rag.index_status(settings, options)
    second.unlink()
    deleted = rag.index_status(settings, options)

    assert len(calls) == 3
    assert initial["documents_processed"] == 2
    assert updated["documents_processed"] == 1
    assert updated["documents_reused"] == 1
    assert deleted["documents_processed"] == 0
    assert deleted["documents_reused"] == 1
    assert deleted["documents_removed"] == 1
    assert deleted["llm_calls"] == 0


def test_an_unchanged_index_is_reused_with_zero_work(rag_env):
    rag_env.write("knowledge", "a.md", "alpha document")
    settings = rag_env.settings()

    rag.index_status(settings)
    again = rag.index_status(settings)

    assert again["index_reused"] is True
    assert again["documents_reused"] == 1 and again["llm_calls"] == 0 and again["context_calls"] == 0


def test_a_stale_index_is_rebuilt_the_way_it_was_built(rag_env, monkeypatch):
    note = rag_env.write("knowledge", "a.md", "alpha document")
    seen = []
    real = pipeline.split_document

    def spy(text, **kwargs):
        seen.append(kwargs["mode"])
        return real(text, **{**kwargs, "mode": "simple"})

    monkeypatch.setattr(pipeline, "split_document", spy)
    settings = rag_env.settings()
    rag.index_status(settings, BuildOptions(split_mode="llm", chunk_model="m"))
    note.write_text("alpha document changed", encoding="utf-8")

    rag.search("alpha", settings)  # no options: the existing index decides how to rebuild

    assert seen == ["llm", "llm"]


def test_rag_ingests_uploaded_docx_as_markdown(rag_env):
    from docx import Document

    upload = rag_env.root / "backend" / "runtime" / "uploads" / ("a" * 32) / "report.docx"
    upload.parent.mkdir(parents=True)
    document = Document()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("Revenue increased in the current quarter.")
    document.save(upload)
    settings = rag_env.settings()

    result = rag.ingest_uploaded_file(
        RagRequest(action="ingest", path="backend/runtime/uploads/" + "a" * 32 + "/report.docx", name="quarterly-report.md"),
        settings,
        BuildOptions(),
    )

    saved = rag_env.root / "user_knowledge" / "quarterly-report.md"
    assert result["path"].endswith("user_knowledge/quarterly-report.md")
    assert result["format"] == "docx"
    assert "Imported from `report.docx` using fileReader" in saved.read_text(encoding="utf-8")
    assert result["index"]["split_mode"] == "simple"
    assert "Revenue increased" in "\n".join(rag.search("revenue quarter", settings))


def test_the_chunk_cache_survives_a_partial_write(rag_env):
    rag_env.write("knowledge", "a.md", "alpha")
    settings = rag_env.settings()
    rag.index_status(settings)
    cache = chunk_cache.cache_path(settings)

    cache.write_text('{"version": 1, "documents": {', encoding="utf-8")  # a torn file is just a cache miss

    assert chunk_cache.load_cache(settings)["documents"] == {}
    assert rag.index_status(settings, force=True)["chunk_count"] == 1
    assert paths.relative_path(cache).endswith("index.chunks.json")
