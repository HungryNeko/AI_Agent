"""Read common local document formats without modifying them."""

from __future__ import annotations

import json
import mimetypes
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from tools.settings import FileReaderSettings

TEXT_SUFFIXES = {
    ".c",
    ".cfg",
    ".conf",
    ".cpp",
    ".css",
    ".csv",
    ".go",
    ".h",
    ".hpp",
    ".htm",
    ".html",
    ".ini",
    ".ipynb",
    ".java",
    ".js",
    ".json",
    ".jsonl",
    ".jsx",
    ".log",
    ".md",
    ".py",
    ".ps1",
    ".rst",
    ".rs",
    ".scss",
    ".sh",
    ".sql",
    ".tex",
    ".toml",
    ".ts",
    ".tsv",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx"}
LEGACY_OFFICE_SUFFIXES = {".doc", ".ppt", ".xls"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | DOCUMENT_SUFFIXES
PROTECTED_PARTS = {".git", "node_modules", "__pycache__", ".venv", ".conda"}
PROTECTED_NAMES = {
    "api-key.txt",
    "api_configs.local.json",
    "servers.local.json",
    "settings.local.json",
}
MAX_ARCHIVE_MEMBERS = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100_000_000


@dataclass(frozen=True)
class FileReadRequest:
    path: str
    start_page: int | None = None
    end_page: int | None = None
    sheet: str = ""
    max_chars: int | None = None


def read(request: FileReadRequest, settings: FileReaderSettings) -> dict[str, Any]:
    """Extract bounded text and metadata from a supported project or uploaded file."""

    path = resolve_path(request.path, settings)
    size = path.stat().st_size
    if size > settings.max_file_bytes:
        raise ValueError(f"file is too large to read: {size} bytes")

    suffix = path.suffix.lower()
    if suffix in LEGACY_OFFICE_SUFFIXES:
        raise ValueError(
            f"legacy {suffix} files are not supported; convert the file to "
            f"{modern_office_suffix(suffix)} first."
        )
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ValueError(f"unsupported file type {suffix or '(no extension)'}; supported: {supported}")

    limit = requested_limit(request.max_chars, settings.max_output_chars)
    content, metadata = extract_content(path, request, soft_limit=limit * 2)
    content = content.strip()
    extracted_characters = len(content)
    truncated = len(content) > limit
    if truncated:
        content = content[:limit].rstrip() + "\n...[content truncated]"

    return {
        "path": relative_path(path, resolve_root(settings)),
        "format": suffix.removeprefix("."),
        "contentType": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "sizeBytes": size,
        "extractedCharacters": extracted_characters,
        "returnedCharacters": len(content),
        "truncated": truncated,
        **metadata,
        "content": content,
    }


def extract_content(
    path: Path,
    request: FileReadRequest,
    *,
    soft_limit: int,
) -> tuple[str, dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, request, soft_limit=soft_limit)
    if suffix == ".docx":
        ensure_safe_office_archive(path)
        return extract_docx(path, soft_limit=soft_limit)
    if suffix == ".pptx":
        ensure_safe_office_archive(path)
        return extract_pptx(path, request, soft_limit=soft_limit)
    if suffix == ".xlsx":
        ensure_safe_office_archive(path)
        return extract_xlsx(path, request, soft_limit=soft_limit)
    if suffix in {".html", ".htm"}:
        return extract_html(path), {}
    return extract_text(path, suffix), {}


def extract_pdf(
    path: Path,
    request: FileReadRequest,
    *,
    soft_limit: int,
) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency error is environment-specific
        raise RuntimeError("PDF reading requires the `pypdf` package.") from exc

    reader = PdfReader(str(path))
    total = len(reader.pages)
    start, end = page_range(request.start_page, request.end_page, total)
    sections: list[str] = []
    used = 0
    for page_number in range(start, end + 1):
        text = reader.pages[page_number - 1].extract_text() or ""
        section = f"[Page {page_number}]\n{text.strip()}"
        sections.append(section)
        used += len(section)
        if used >= soft_limit:
            break
    return "\n\n".join(sections), {
        "pageCount": total,
        "startPage": start,
        "endPage": min(end, start + len(sections) - 1),
        "needsOcr": not any(section.split("\n", 1)[-1].strip() for section in sections),
    }


def extract_docx(path: Path, *, soft_limit: int) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Word reading requires the `python-docx` package.") from exc

    document = Document(str(path))
    parts: list[str] = []
    used = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
            used += len(text)
        if used >= soft_limit:
            break
    table_count = 0
    if used < soft_limit:
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append("\t".join(clean_cell(cell.text) for cell in row.cells))
            table_text = f"[Table {table_index}]\n" + "\n".join(rows)
            parts.append(table_text)
            used += len(table_text)
            table_count += 1
            if used >= soft_limit:
                break
    return "\n\n".join(parts), {
        "paragraphCount": len(document.paragraphs),
        "tableCount": len(document.tables),
        "tablesExtracted": table_count,
    }


def extract_pptx(
    path: Path,
    request: FileReadRequest,
    *,
    soft_limit: int,
) -> tuple[str, dict[str, Any]]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("PowerPoint reading requires the `python-pptx` package.") from exc

    presentation = Presentation(str(path))
    total = len(presentation.slides)
    start, end = page_range(request.start_page, request.end_page, total)
    sections: list[str] = []
    used = 0
    for slide_number in range(start, end + 1):
        slide = presentation.slides[slide_number - 1]
        items: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = str(getattr(shape, "text", "")).strip()
                if text:
                    items.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    items.append("\t".join(clean_cell(cell.text) for cell in row.cells))
        section = f"[Slide {slide_number}]\n" + "\n".join(items)
        sections.append(section)
        used += len(section)
        if used >= soft_limit:
            break
    return "\n\n".join(sections), {
        "slideCount": total,
        "startPage": start,
        "endPage": min(end, start + len(sections) - 1),
    }


def extract_xlsx(
    path: Path,
    request: FileReadRequest,
    *,
    soft_limit: int,
) -> tuple[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel reading requires the `openpyxl` package.") from exc

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if request.sheet:
            if request.sheet not in workbook.sheetnames:
                raise ValueError(
                    f"worksheet not found: {request.sheet}; available: {', '.join(workbook.sheetnames)}"
                )
            worksheets = [workbook[request.sheet]]
        else:
            worksheets = list(workbook.worksheets)

        sections: list[str] = []
        used = 0
        sheets_read: list[str] = []
        for worksheet in worksheets:
            lines = [f"[Sheet: {worksheet.title}]"]
            for row in worksheet.iter_rows(values_only=True):
                values = [format_cell(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    lines.append("\t".join(values))
                if sum(len(line) for line in lines) + used >= soft_limit:
                    break
            section = "\n".join(lines)
            sections.append(section)
            sheets_read.append(worksheet.title)
            used += len(section)
            if used >= soft_limit:
                break
        return "\n\n".join(sections), {
            "sheetNames": workbook.sheetnames,
            "sheetsRead": sheets_read,
        }
    finally:
        workbook.close()


def extract_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("HTML reading requires the `beautifulsoup4` package.") from exc
    html = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript"]):
        element.decompose()
    return "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())


def extract_text(path: Path, suffix: str) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".ipynb":
        return extract_notebook(text)
    return text


def extract_notebook(text: str) -> str:
    try:
        notebook = json.loads(text)
    except json.JSONDecodeError:
        return text
    sections: list[str] = []
    for index, cell in enumerate(notebook.get("cells", []), start=1):
        if not isinstance(cell, dict):
            continue
        source = cell.get("source", [])
        content = "".join(source) if isinstance(source, list) else str(source or "")
        if content.strip():
            sections.append(f"[Cell {index}: {cell.get('cell_type', 'unknown')}]\n{content.strip()}")
    return "\n\n".join(sections)


def ensure_safe_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise ValueError("Office document contains too many archive members.")
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError("Office document expands beyond the safe extraction limit.")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid or corrupted Office document.") from exc


def page_range(start_page: int | None, end_page: int | None, total: int) -> tuple[int, int]:
    if total <= 0:
        return 1, 0
    start = 1 if start_page is None else start_page
    end = total if end_page is None else end_page
    if start < 1 or end < 1:
        raise ValueError("startPage and endPage must be positive 1-based numbers.")
    if end < start:
        raise ValueError("endPage must be greater than or equal to startPage.")
    if start > total:
        raise ValueError(f"startPage {start} exceeds document page count {total}.")
    return start, min(end, total)


def requested_limit(value: int | None, maximum: int) -> int:
    if value is None:
        return maximum
    return max(1_000, min(int(value), maximum))


def resolve_root(settings: FileReaderSettings) -> Path:
    if settings.root:
        root = Path(settings.root)
        if not root.is_absolute():
            root = project_root() / root
        return root.resolve()
    return project_root()


def resolve_path(path_text: str, settings: FileReaderSettings) -> Path:
    clean = normalize_upload_reference(path_text)
    if not clean:
        raise ValueError("fileReader path is required.")
    root = resolve_root(settings)
    raw = Path(clean)
    path = raw if raw.is_absolute() else root / raw
    resolved = path.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise ValueError("path must stay inside the configured file reader root.") from exc
    if not resolved.is_file():
        raise ValueError(f"path is not a file: {path_text}")
    ensure_not_protected(resolved, root)
    return resolved


def normalize_upload_reference(value: str) -> str:
    clean = value.strip()
    parsed = urlparse(clean)
    if parsed.scheme in {"http", "https"}:
        clean = parsed.path
    clean = unquote(clean.split("?", 1)[0])
    if clean.startswith("/api/uploads/"):
        parts = clean.split("/", 4)
        if len(parts) != 5 or not parts[3] or not parts[4]:
            raise ValueError("upload URL must be /api/uploads/{upload_id}/{filename}.")
        return f"backend/runtime/uploads/{parts[3]}/{parts[4]}"
    if parsed.scheme:
        raise ValueError("fileReader accepts local project paths or local /api/uploads URLs only.")
    return clean


def ensure_not_protected(path: Path, root: Path) -> None:
    rel = relative_path(path, root).replace("\\", "/")
    parts = set(rel.split("/"))
    lower_name = path.name.lower()
    if lower_name == ".env" or lower_name.startswith(".env.") or lower_name in PROTECTED_NAMES:
        raise ValueError("protected or secret-bearing file cannot be read.")
    if any(part in parts for part in PROTECTED_PARTS):
        raise ValueError("protected path cannot be read.")
    if rel.startswith("backend/runtime/") and not rel.startswith("backend/runtime/uploads/"):
        raise ValueError("only uploaded files may be read under backend/runtime.")


def modern_office_suffix(suffix: str) -> str:
    return {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx"}[suffix]


def clean_cell(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.replace("\r", " ").replace("\n", " ").split())


def format_cell(value: object) -> str:
    if value is None:
        return ""
    return clean_cell(value)


def relative_path(path: Path, root: Path) -> str:
    try:
        return str(path.relative_to(root)).replace("\\", "/")
    except ValueError:
        return str(path)


def project_root() -> Path:
    return Path(__file__).resolve().parents[2]
